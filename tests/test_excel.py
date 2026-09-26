import tempfile
import unittest
from pathlib import Path
import openpyxl

from app.reports.excel import ExcelReportGenerator
from app.matching.schema import (
    CanonicalJob,
    JobMatch,
    MatchCategory,
    ScoreBreakdown,
    PipelineRunStats,
    ApplicationStatus
)


class TestExcelReportGenerator(unittest.TestCase):
    def test_excel_generation_structure(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = Path(tmpdir)
            gen = ExcelReportGenerator(output_dir=out_path)

            job = CanonicalJob(
                canonical_id="job-123",
                title="Data Analyst",
                company="RetailPulse",
                location="Bangalore",
                work_mode="Remote",
                description="SQL & Power BI",
                experience_required="1-2 yrs",
                education_required="B.Sc.",
                application_url="https://example.com/apply/123",
                sources=["LinkedIn", "Naukri"]
            )
            match = JobMatch(
                job=job,
                relevance_score=92.5,
                category=MatchCategory.HIGH_MATCH,
                score_breakdown=ScoreBreakdown(total_relevance_score=92.5),
                matched_skills=["SQL", "Python"],
                why_matches=["Strong SQL alignment"],
                application_status=ApplicationStatus.NOT_APPLIED
            )
            stats = PipelineRunStats(
                run_id="test_run",
                jobs_discovered=10,
                jobs_deduplicated=8,
                jobs_matched=8,
                jobs_recommended=7,
                high_matches=1
            )

            file_created = gen.generate_report([match], stats)
            self.assertTrue(Path(file_created).exists())

            # Verify with openpyxl
            wb = openpyxl.load_workbook(file_created)
            self.assertIn("Summary", wb.sheetnames)
            self.assertIn("Job Matches", wb.sheetnames)

            ws_jobs = wb["Job Matches"]
            self.assertEqual(ws_jobs.cell(row=1, column=1).value, "Match Score")
            self.assertEqual(ws_jobs.cell(row=2, column=3).value, "Data Analyst")
            self.assertEqual(ws_jobs.cell(row=2, column=4).value, "RetailPulse")
            self.assertEqual(ws_jobs.cell(row=2, column=13).value, "LinkedIn, Naukri")


if __name__ == "__main__":
    unittest.main()
