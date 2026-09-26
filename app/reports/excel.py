import logging
from datetime import datetime
from pathlib import Path
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import settings
from app.matching.schema import JobMatch, MatchCategory, PipelineRunStats

logger = logging.getLogger(__name__)


class ExcelReportGenerator:
    """Generates styled multi-sheet Excel reports with summary metrics and ranked jobs."""

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or settings.get_absolute_report_dir()

    def generate_report(self, matches: List[JobMatch], stats: PipelineRunStats) -> str:
        date_str = datetime.utcnow().strftime("%Y-%m-%d")
        filename = f"job_matches_{date_str}.xlsx"
        file_path = self.output_dir / filename

        wb = openpyxl.Workbook()

        # Styles
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
        subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        regular_font = Font(name="Segoe UI", size=10)
        link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        category_fills = {
            MatchCategory.HIGH_MATCH.value: PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
            MatchCategory.GOOD_MATCH.value: PatternFill(start_color="D0E0E3", end_color="D0E0E3", fill_type="solid"),
            MatchCategory.STRETCH.value: PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),
            MatchCategory.LOW_MATCH.value: PatternFill(start_color="F3F3F3", end_color="F3F3F3", fill_type="solid"),
        }

        # -------------------------------------------------------------
        # SHEET 1: Summary Sheet
        # -------------------------------------------------------------
        ws_summary = wb.active
        ws_summary.title = "Summary"
        ws_summary.views.sheetView[0].showGridLines = True

        ws_summary["A1"] = "Job Discovery & Relevance Matching Summary"
        ws_summary["A1"].font = title_font
        ws_summary["A2"] = f"Generated on {date_str} (UTC) | Automated Run ID: {stats.run_id}"
        ws_summary["A2"].font = subtitle_font

        # Key Metrics Table
        headers_metrics = ["Metric", "Count"]
        ws_summary.cell(row=4, column=1, value="Key Performance Metric").font = bold_font
        ws_summary.cell(row=4, column=2, value="Value").font = bold_font

        metrics_data = [
            ("Report Date", date_str),
            ("Jobs Discovered", stats.jobs_discovered),
            ("Jobs After Deduplication", stats.jobs_deduplicated),
            ("Jobs Evaluated", stats.jobs_matched),
            ("High Matches (>=85%)", stats.high_matches),
            ("Good Matches (>=70%)", stats.good_matches),
            ("Stretch Matches (>=55%)", stats.stretch_matches),
            ("Recommended for Consideration", stats.jobs_recommended),
        ]

        curr_row = 5
        for label, val in metrics_data:
            c1 = ws_summary.cell(row=curr_row, column=1, value=label)
            c2 = ws_summary.cell(row=curr_row, column=2, value=val)
            c1.font = regular_font
            c2.font = bold_font if "Matches" in label or "Recommended" in label else regular_font
            c1.border = thin_border
            c2.border = thin_border
            curr_row += 1

        # Breakdown by Source
        curr_row += 2
        ws_summary.cell(row=curr_row, column=1, value="Breakdown by Source").font = bold_font
        ws_summary.cell(row=curr_row, column=2, value="Listings").font = bold_font
        curr_row += 1

        source_counts = {}
        for m in matches:
            for s in m.job.sources:
                source_counts[s] = source_counts.get(s, 0) + 1

        for src, count in sorted(source_counts.items(), key=lambda x: x[1], reverse=True):
            c1 = ws_summary.cell(row=curr_row, column=1, value=src)
            c2 = ws_summary.cell(row=curr_row, column=2, value=count)
            c1.font = regular_font
            c2.font = regular_font
            c1.border = thin_border
            c2.border = thin_border
            curr_row += 1

        # Breakdown by Location
        curr_row += 2
        ws_summary.cell(row=curr_row, column=1, value="Breakdown by Location").font = bold_font
        ws_summary.cell(row=curr_row, column=2, value="Jobs").font = bold_font
        curr_row += 1

        loc_counts = {}
        for m in matches:
            loc = m.job.location or "Unspecified"
            loc_counts[loc] = loc_counts.get(loc, 0) + 1

        for loc, count in sorted(loc_counts.items(), key=lambda x: x[1], reverse=True):
            c1 = ws_summary.cell(row=curr_row, column=1, value=loc)
            c2 = ws_summary.cell(row=curr_row, column=2, value=count)
            c1.font = regular_font
            c2.font = regular_font
            c1.border = thin_border
            c2.border = thin_border
            curr_row += 1

        ws_summary.column_dimensions["A"].width = 35
        ws_summary.column_dimensions["B"].width = 25

        # -------------------------------------------------------------
        # SHEET 2: Job Matches
        # -------------------------------------------------------------
        ws_jobs = wb.create_sheet(title="Job Matches")
        ws_jobs.views.sheetView[0].showGridLines = True

        columns = [
            ("Match Score", 14),
            ("Category", 15),
            ("Job Title", 28),
            ("Company", 24),
            ("Location", 20),
            ("Work Mode", 14),
            ("Experience Required", 20),
            ("Salary", 22),
            ("Matched Skills", 32),
            ("Missing Skills", 28),
            ("Why This Matches", 45),
            ("Posted Date", 16),
            ("Source", 18),
            ("Application URL", 24),
            ("Application Status", 18),
        ]

        # Write Headers
        for col_idx, (col_name, _) in enumerate(columns, start=1):
            cell = ws_jobs.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        ws_jobs.row_dimensions[1].height = 28

        # Populate rows
        for row_idx, match in enumerate(matches, start=2):
            score_val = f"{match.relevance_score:.1f}%"
            category_val = match.category.value
            title_val = match.job.title
            company_val = match.job.company
            loc_val = match.job.location
            work_mode_val = match.job.work_mode
            exp_val = match.job.experience_required or "Not specified"
            salary_val = match.job.salary or "Not disclosed"
            matched_skills_str = ", ".join(match.matched_skills) if match.matched_skills else "None explicitly listed"
            missing_skills_str = ", ".join(match.missing_skills) if match.missing_skills else "None"
            why_str = " \n• ".join(["• " + r for r in match.why_matches]) if match.why_matches else ""
            posted_date_str = match.job.posted_at.strftime("%Y-%m-%d") if match.job.posted_at else "Unknown"
            sources_str = ", ".join(match.job.sources) if match.job.sources else "Direct"
            app_url = match.job.application_url
            app_status = match.application_status.value

            row_data = [
                score_val,
                category_val,
                title_val,
                company_val,
                loc_val,
                work_mode_val,
                exp_val,
                salary_val,
                matched_skills_str,
                missing_skills_str,
                why_str,
                posted_date_str,
                sources_str,
                "Apply Here",
                app_status,
            ]

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws_jobs.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border

                if col_idx == 1:
                    cell.value = val
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                    cell.font = bold_font
                elif col_idx == 2:
                    cell.value = val
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                    if val in category_fills:
                        cell.fill = category_fills[val]
                elif col_idx == 14:
                    cell.value = "Apply Link"
                    if app_url and app_url.startswith("http"):
                        cell.hyperlink = app_url
                        cell.font = link_font
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif col_idx in (9, 10, 11):
                    cell.value = val
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    cell.value = val
                    cell.alignment = Alignment(horizontal="left", vertical="top")

            ws_jobs.row_dimensions[row_idx].height = 45

        # Freeze top row and set autofilter
        ws_jobs.freeze_panes = "A2"
        max_col_letter = get_column_letter(len(columns))
        ws_jobs.auto_filter.ref = f"A1:{max_col_letter}{len(matches) + 1}"

        # Adjust column widths
        for col_idx, (_, col_width) in enumerate(columns, start=1):
            letter = get_column_letter(col_idx)
            ws_jobs.column_dimensions[letter].width = col_width

        try:
            wb.save(str(file_path))
            logger.info(f"Excel report successfully written to: {file_path}")
            return str(file_path)
        except PermissionError:
            timestamp = datetime.utcnow().strftime("%H%M%S")
            fallback_filename = f"job_matches_{date_str}_{timestamp}.xlsx"
            fallback_path = self.output_dir / fallback_filename
            logger.warning(
                f"File '{file_path}' is locked by another program (e.g., open in Microsoft Excel). "
                f"Saved fallback report to '{fallback_path}' instead."
            )
            wb.save(str(fallback_path))
            return str(fallback_path)
